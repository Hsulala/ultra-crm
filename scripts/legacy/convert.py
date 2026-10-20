"""把客戶舊的 Google Sheets（課堂報表、每月銷售總表）轉成匯入用的 JSON。

只在本機跑（需要 openpyxl，正式環境用不到所以沒放進 requirements.txt）：

    pip install openpyxl
    python -m scripts.legacy.convert --contracts 課堂報表.xlsx \
        --revenue-dir 銷售總表資料夾/ --out bundle.json

輸出的 bundle.json 交給後台「匯入舊資料」（app/services/legacy_import.py）處理。
這支腳本只讀檔案、不碰資料庫；看不懂的儲存格不會猜，原樣放進 issues 清單
讓人確認。資料檔本身含客人姓名，不要加進 git。"""

import argparse
import collections
import datetime as dt
import glob
import json
import os
import re
import warnings

warnings.filterwarnings("ignore")

# 店內正職教練（客戶銷售總表裡只有這 6 位會開「教練課」，其他人都是付場租的場租教練）
STAFF_COACHES = ["NICK", "WEILIN", "YOKO", "DK", "NINA", "MUMU"]

# 明顯的拼字錯誤，統一成同一位教練
COACH_ALIASES = {
    "ELAS": "ELSA",
    "SABRIAN": "SABRINA",
    "CHIRS": "CHRIS",
    "JAEVN": "JAVEN",
    "JAVAN": "JAVEN",
    "LARY": "LARRY",
    "BRAIN": "BRIAN",
    "JEERY": "JERRY",
    "MILLS": "MILES",
    "陳建鋐": "建鋐",
    "陳龍毅": "龍毅",
    "EIDDIE": "EDDIE",
    "JAGUR": "JAGUAR",
    "JAGUW": "JAGUAR",
    "MILIS": "MILES",
    "胡偉翔": "胡緯翔",
    "琪晶": "晶琪",
}

SESSION_COLS = range(7, 32)  # H..AF：第 1～25 堂
NOTE_WORDS = ("退費", "遺失", "終止", "換約", "轉約", "轉換", "更改", "補差額", "改1v", "改1V")
REFUND_WORDS = ("退費",)


def norm_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def norm_coach(v):
    raw = norm_text(v).upper().replace(" ", "")
    if not raw or raw == "NONE":
        return "", raw
    parts = [COACH_ALIASES.get(p, p) for p in re.split(r"[/、,]", raw) if p]
    return (parts[0] if parts else ""), raw


def _valid(y, m, d):
    try:
        out = dt.date(y, m, d)
    except ValueError:
        return None
    if not (dt.date(2019, 1, 1) <= out <= dt.date(2027, 12, 31)):
        return None
    return out


def parse_date(v):
    """回傳 (date 或 None, 是否看起來像日期)。"""
    if isinstance(v, (dt.datetime, dt.date)):
        d = v.date() if isinstance(v, dt.datetime) else v
        return (_valid(d.year, d.month, d.day), True)
    t = norm_text(v)
    if not t:
        return (None, False)
    t = t.replace("．", ".").replace("-", "/").replace(" ", "")
    t = re.sub(r"\.{2,}", ".", t).strip(".")
    # 西元 2024/11/23；常見打錯：0205/5/30 → 2025、202412/27
    m = re.fullmatch(r"(\d{4})/(\d{1,2})/(\d{1,2})", t)
    if m:
        y, mo, d = map(int, m.groups())
        if y == 205:
            y = 2025
        elif y < 1000:  # 0113/10/1 這種其實是民國
            y += 1911
        return (_valid(y, mo, d), True)
    m = re.fullmatch(r"(\d{4})(\d{2})/(\d{1,2})", t)
    if m:
        return (_valid(int(m.group(1)), int(m.group(2)), int(m.group(3))), True)
    # 民國 110.8.4；常見打錯：少打開頭的 1（10.8.24 → 110）
    m = re.fullmatch(r"(\d{2,3})[./](\d{1,2})[./](\d{1,2})", t)
    if m:
        y, mo, d = map(int, m.groups())
        if y < 100:
            y += 100
        return (_valid(y + 1911, mo, d), True)
    looks = bool(re.fullmatch(r"[\d./]+", t))
    return (None, looks)


def read_contract_sheet(ws, course_type, issues):
    out = []
    for ln, r in enumerate(ws.iter_rows(values_only=True), start=1):
        r = list(r) + [None] * (32 - len(r))
        no = norm_text(r[1])
        if not no or no == "合約編號":
            continue
        where = f"{ws.title} 第 {ln} 列"
        if not re.fullmatch(r"\d+", no):
            if re.search(r"\d", no):
                issues.append({"where": where, "kind": "合約編號看不懂", "value": no})
            continue
        contract_no = no.zfill(5)
        names = [n.strip() for n in re.split(r"[/／]", norm_text(r[2])) if n.strip()]
        if not names:
            issues.append({"where": where, "kind": "沒有學員姓名，略過", "value": contract_no})
            continue
        coach, coach_raw = norm_coach(r[3])
        notes = []
        if coach_raw and "/" in coach_raw:
            notes.append(f"原表教練：{coach_raw}")

        status = "active"
        sessions = []
        for idx in SESSION_COLS:
            v = r[idx]
            t = norm_text(v)
            if t == "" or t.upper().rstrip(".") == "X":
                continue
            d, looks = parse_date(v)
            if d is not None:
                sessions.append({"no": idx - 6, "date": d.isoformat()})
            elif looks:
                sessions.append({"no": idx - 6, "date": None, "raw": t})
                issues.append({"where": where, "kind": "上課日期打錯，已記為上過一堂、日期留原文", "value": f"{contract_no} 第 {idx-6} 堂：{t}"})
            elif any(w in t for w in NOTE_WORDS):
                notes.append(t)
                if any(w in t for w in REFUND_WORDS):
                    status = "refunded"
            elif t in ("#VALUE!", "00:00:00", "\\", "/"):
                issues.append({"where": where, "kind": "儲存格是錯誤值，略過", "value": f"{contract_no} 第 {idx-6} 堂：{t}"})
            else:
                notes.append(t)
                issues.append({"where": where, "kind": "看不懂的上課欄位，放進備註", "value": f"{contract_no}：{t}"})

        purchase, _ = parse_date(r[4])
        ptext = norm_text(r[4])
        if purchase is None:
            dated = [s["date"] for s in sessions if s["date"]]
            if ptext and not re.fullmatch(r"[\d./]+", ptext):
                notes.append(f"原表購買日期欄：{ptext}")
            if dated:
                purchase = dt.date.fromisoformat(min(dated))
                issues.append({"where": where, "kind": "購買日期看不懂，改用第一堂日期", "value": f"{contract_no}：{ptext or '空白'} → {purchase}"})
            else:
                issues.append({"where": where, "kind": "沒有購買日期也沒有上課日期，略過", "value": f"{contract_no}：{ptext or '空白'}"})
                continue

        ptxt = norm_text(r[5])
        mnum = re.match(r"\d+(\.\d+)?", ptxt)
        if ptxt.upper().startswith("FREE"):
            unit_price = 0.0
            if ptxt.upper() != "FREE":
                notes.append(f"原表單價：{ptxt}")
            if status == "active":
                status = "free"
        elif mnum:
            unit_price = float(mnum.group(0))
            if mnum.group(0) != ptxt:
                notes.append(f"原表單價：{ptxt}")
        else:
            unit_price = 0.0
            if ptxt:
                notes.append(f"原表單價：{ptxt}")
                issues.append({"where": where, "kind": "單價看不懂，記為 0", "value": f"{contract_no}：{ptxt}"})
        try:
            total = int(float(r[6]))
        except (TypeError, ValueError):
            total = max(len(sessions), 1)
            issues.append({"where": where, "kind": "堂數看不懂，改用已上堂數", "value": f"{contract_no}：{r[6]} → {total}"})
        if len(sessions) > total:
            issues.append({"where": where, "kind": "已上堂數比合約堂數多", "value": f"{contract_no}：{len(sessions)}/{total}"})
        if not coach:
            issues.append({"where": where, "kind": "沒有教練", "value": contract_no})

        out.append(
            {
                "sheet_line": ln,
                "contract_no": contract_no,
                "course_type": course_type,
                "students": names,
                "coach": coach,
                "purchase_date": purchase.isoformat(),
                "unit_price": unit_price,
                "total_sessions": total,
                "status": status,
                "note": "；".join(notes) or None,
                "sessions": sessions,
            }
        )
    return out


REV_KEYS = {"日期": "date", "品項": "cat", "姓名": "name", "教練": "coach", "內容": "content", "付款方式": "pay", "備註": "note"}
CAT_MAP = {"飲料": "beverage", "教練課": "coaching", "會籍": "membership", "服務": "service"}


def read_revenue_file(path, issues):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = [w for w in wb.worksheets if "手動輸入" in w.title]
    label = os.path.basename(path)
    if not sheets:
        issues.append({"where": label, "kind": "找不到「(手動輸入)銷售總表」分頁", "value": ""})
        return []
    rows = list(sheets[0].iter_rows(values_only=True))
    col = {}
    for i, h in enumerate(rows[0]):
        h = norm_text(h)
        if "金額" in h:
            col.setdefault("amt", i)
        elif h in REV_KEYS:
            col.setdefault(REV_KEYS[h], i)
    out = []
    for ln, r in enumerate(rows[1:], start=2):
        g = lambda k: r[col[k]] if k in col and col[k] < len(r) else None  # noqa: E731
        d, amt, content = g("date"), g("amt"), norm_text(g("content"))
        if d in (None, "") and amt in (None, "") and not content:
            continue
        where = f"{label} 第 {ln} 列"
        day, _ = parse_date(d)
        if day is None or not isinstance(amt, (int, float)):
            issues.append({"where": where, "kind": "營收列缺日期或金額，略過", "value": f"{d}｜{content}｜{amt}"})
            continue
        cat_raw = norm_text(g("cat"))
        cat = CAT_MAP.get(cat_raw, "custom")
        if cat == "service" and "場租" in content:
            cat = "venue_rental"
        coach, _ = norm_coach(g("coach"))
        pay = norm_text(g("pay")) or "未填"
        note = norm_text(g("note"))
        out.append(
            {
                "source": f"{label}#{ln}",
                "entry_date": day.isoformat(),
                "item_category": cat,
                "item_name": content or cat_raw or "未填",
                "customer_name": norm_text(g("name")) or None,
                "coach": coach or None,
                "payment_method": pay,
                "amount": float(amt),
                "note": note or None,
            }
        )
    return out


def main():
    import openpyxl

    ap = argparse.ArgumentParser()
    ap.add_argument("--contracts", required=True)
    ap.add_argument("--revenue-dir")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    issues = []
    wb = openpyxl.load_workbook(a.contracts, data_only=True)
    contracts = []
    for ws in wb.worksheets:
        ctype = "pilates" if "皮拉" in ws.title else "strength"
        contracts += read_contract_sheet(ws, ctype, issues)

    revenue = []
    if a.revenue_dir:
        files = sorted(glob.glob(os.path.join(a.revenue_dir, "*銷售總表*.xlsx")))
        for f in files:
            revenue += read_revenue_file(f, issues)

    coaches = collections.Counter(c["coach"] for c in contracts if c["coach"])
    coaches.update(r["coach"] for r in revenue if r["coach"])
    staff = [c for c in STAFF_COACHES]
    venue = sorted(c for c in coaches if c not in STAFF_COACHES)

    bundle = {
        "format": "ultra-legacy-v1",
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "staff_coaches": staff,
        "venue_coaches": venue,
        "contracts": contracts,
        "revenue": revenue,
        "issues": issues,
    }
    with open(a.out, "w", encoding="utf-8") as fh:
        json.dump(bundle, fh, ensure_ascii=False)

    used = sum(len(c["sessions"]) for c in contracts)
    print(f"合約 {len(contracts)} 筆（重訓 {sum(c['course_type']=='strength' for c in contracts)}、"
          f"皮拉提斯 {sum(c['course_type']=='pilates' for c in contracts)}），上課紀錄 {used} 堂")
    print(f"營收 {len(revenue)} 筆，合計 {sum(r['amount'] for r in revenue):,.0f} 元")
    print(f"正職教練 {len(staff)} 位、場租教練 {len(venue)} 位；需要確認的地方 {len(issues)} 筆")


if __name__ == "__main__":
    main()
